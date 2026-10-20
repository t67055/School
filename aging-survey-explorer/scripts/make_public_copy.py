"""Make the public copy of the workbook: drops the group roster sheet and clears interviewer names.
Usage: python3 make_public_copy.py tagged.xlsx ../data/Project_1_combined_dataset.xlsx"""
import sys
import openpyxl
from openpyxl.cell.cell import MergedCell

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
if 'Member_Groups' in wb.sheetnames:
    del wb['Member_Groups']

def clear_column(ws, header):
    heads = [c.value for c in ws[1]]
    if header not in heads:
        return
    col = heads.index(header) + 1
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, col)
        if not isinstance(c, MergedCell):
            c.value = None

clear_column(wb['Combined'], 'Interviewer')
if 'Group 7' in wb.sheetnames:
    clear_column(wb['Group 7'], 'Interviewer / Group Member Name (Internal Tracking)')
for row in wb['Notes'].iter_rows():
    for c in row:
        if not isinstance(c, MergedCell) and isinstance(c.value, str) and 'Member_Groups' in c.value:
            c.value = '• Interviewer names and the group roster are removed from this public copy.'
wb.save(out)
print('saved', out)
