import sys, json, collections
sys.path.insert(0, '.')
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo
from tagger import tag, action_status, GENERAL, DEFINITION, MAINTAIN, ACTIONS, REASON_EXTRA

SRC, OUT = sys.argv[1], sys.argv[2]
df = pd.read_pickle('combined.pkl')

# (short key, source column, short label, codebook or None for Q9 status)
QUESTIONS = [
 ('q5',  'Q5 How would you define aging?', 'Q5 Definition of aging', DEFINITION),
 ('q7',  'Q7 Greatest challenge facing the aging population', 'Q7 Greatest challenge', GENERAL),
 ('q8',  'Q8 Challenge to solve in the short term', 'Q8 Short-term priority', GENERAL),
 ('q9',  'Q9 Action being taken on this challenge?', 'Q9 Is action being taken?', None),
 ('q10', 'Q10 Three things to maintain through life', 'Q10 Things to maintain', MAINTAIN),
 ('q11', 'Q11 What you can do to maintain them', 'Q11 What you can do', ACTIONS),
 ('rank', 'Reason for ranking', 'Reason for topic ranking', REASON_EXTRA + GENERAL),
 ('excite', 'What excites you', 'What excites you', GENERAL),
 ('concern', 'Greatest concern', 'Greatest concern', GENERAL),
 ('ethics', 'Ethical considerations (details)', 'Ethical considerations', GENERAL),
]
TAIL = ['Other', 'General / no specifics', 'No answer / unsure']

tags = {}
for key, col, label, cb in QUESTIONS:
    if cb is None:
        tags[key] = df[col].apply(lambda t: [action_status(t)] if isinstance(t, str) and t.strip() else [])
    else:
        tags[key] = df[col].apply(lambda t: tag(t, cb) if isinstance(t, str) and t.strip() else [])
pd.to_pickle({'tags': tags, 'questions': [(k, c, l) for k, c, l, _ in QUESTIONS]}, 'tags.pkl')

wb = openpyxl.load_workbook(SRC)
FONT = 'Arial'
hdr_fill = PatternFill('solid', fgColor='1F3864')
hdr_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
body = Font(name=FONT, size=10)
bold = Font(name=FONT, bold=True, size=10)

# ---------------------------------------------------------------- Tagged Responses
ws = wb.create_sheet('Tagged Responses', 1)
demo = ['Respondent ID', 'Group', 'Age', 'Age category', 'Gender', 'From – Country', 'Residence – Country']
headers = demo[:]
for key, col, label, cb in QUESTIONS:
    if cb is None:
        headers += [f'{label} – response', f'{label} – status']
    else:
        headers += [f'{label} – response', f'{label} – themes']
ws.append(headers)
comb_cols = {c: i + 1 for i, c in enumerate(df.columns)}
agecat_letter = L(comb_cols['Age category'])
for i, row in df.reset_index(drop=True).iterrows():
    r = i + 2
    vals = [row['Respondent ID'], row['Group'], row['Age'] if pd.notna(row['Age']) else None,
            f'=Combined!{agecat_letter}{r}', row['Gender'], row['From – Country'], row['Residence – Country']]
    for key, col, label, cb in QUESTIONS:
        txt = row[col] if isinstance(row[col], str) else None
        vals += [txt, '; '.join(tags[key].iloc[i]) or None]
    ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in vals])
n = len(df) + 1
widths = [11, 7, 6, 14, 12, 16, 16] + [44, 30] * len(QUESTIONS)
for j, w in enumerate(widths, start=1):
    ws.column_dimensions[L(j)].width = w
    c = ws.cell(1, j)
    c.font, c.fill = hdr_font, hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row:
        c.font = body
        c.alignment = Alignment(vertical='top')
ws.row_dimensions[1].height = 42
ws.freeze_panes = 'B2'
t = Table(displayName='TaggedResponses', ref=f'A1:{L(len(headers))}{n}')
t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
ws.add_table(t)
col_of = {h: L(j) for j, h in enumerate(headers, start=1)}

# ---------------------------------------------------------------- Theme Summary
ss = wb.create_sheet('Theme Summary', 2)
TR = "'Tagged Responses'"
rng = lambda h: f"{TR}!${col_of[h]}$2:${col_of[h]}${n}"
crit = lambda v: v[1:] if v.startswith('=') else f'"{v}"' 
SEGS = [('All respondents', None, None),
        # age-category labels follow the cutoffs on the Notes sheet, so link to its header cells
        ('=Notes!$B$9', 'Age category', '=Notes!$B$9'), ('=Notes!$C$9', 'Age category', '=Notes!$C$9'),
        ('=Notes!$D$9', 'Age category', '=Notes!$D$9'),
        ('Female', 'Gender', 'Female'), ('Male', 'Gender', 'Male'),
        ('Lives in US', 'Residence – Country', 'United States'), ('Lives in China', 'Residence – Country', 'China')]
ss['A1'] = 'Theme summary – share of respondents whose answer mentions each theme'
ss['A1'].font = Font(name=FONT, bold=True, size=13)
ss['A2'] = ('Each answer can carry several themes, so shares in a block can add up to more than 100%. '
            'Percent = respondents in that column whose answer has the theme ÷ respondents in that column who answered the question. '
            'Themes come from keyword codebooks (see the Tagged Responses sheet to check any answer).')
ss['A2'].font = body
ss['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ss.merge_cells('A2:J2')
ss.row_dimensions[2].height = 44
pct_fill = PatternFill('solid', fgColor='EEF2F8')
r = 4
for key, col, label, cb in QUESTIONS:
    resp_h = f'{label} – response'
    tag_h = f'{label} – {"status" if cb is None else "themes"}'
    ss.cell(r, 1, label).font = Font(name=FONT, bold=True, size=11, color='1F3864')
    r += 1
    heads = ['Theme', 'Responses (all)'] + [s[0] for s in SEGS]
    for j, h in enumerate(heads, start=1):
        c = ss.cell(r, j, h)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    ss.row_dimensions[r].height = 30
    r += 1
    # denominators
    ss.cell(r, 1, 'Answered (n)').font = bold
    ss.cell(r, 2, f'=COUNTIF({rng(resp_h)},"?*")').font = bold
    den_row = r
    for j, (name, dcol, dval) in enumerate(SEGS, start=3):
        f = (f'=COUNTIF({rng(resp_h)},"?*")' if dcol is None
             else f'=COUNTIFS({rng(resp_h)},"?*",{rng(dcol)},{crit(dval)})')
        c = ss.cell(r, j, f)
        c.font = bold
    r += 1
    counts = collections.Counter(x for l in tags[key] for x in l)
    order = [k for k, _ in counts.most_common() if k not in TAIL] + [k for k in TAIL if counts.get(k)]
    for theme in order:
        ss.cell(r, 1, theme).font = body
        ss.cell(r, 2, f'=COUNTIF({rng(tag_h)},"*"&$A{r}&"*")').font = body
        for j, (name, dcol, dval) in enumerate(SEGS, start=3):
            num = (f'COUNTIF({rng(tag_h)},"*"&$A{r}&"*")' if dcol is None
                   else f'COUNTIFS({rng(tag_h)},"*"&$A{r}&"*",{rng(dcol)},{crit(dval)})')
            c = ss.cell(r, j, f'=IFERROR({num}/{L(j)}${den_row},"")')
            c.number_format = '0%'
            c.font = body
            c.fill = pct_fill
        r += 1
    r += 1
ss.column_dimensions['A'].width = 40
ss.column_dimensions['B'].width = 11
for j in range(3, 3 + len(SEGS)):
    ss.column_dimensions[L(j)].width = 12
ss.freeze_panes = 'B4'

# Codebook note on the Notes sheet
ns = wb['Notes']
start = ns.max_row + 2
lines = ['Theme tagging (Tagged Responses and Theme Summary sheets)',
 '• Open-ended answers were tagged with keyword codebooks, one per kind of question: a shared set of themes for challenges, priorities, ranking reasons, excitement, concerns and ethics; separate sets for Q5 (definition of aging), Q10 (things to maintain) and Q11 (what you can do).',
 '• An answer can carry several themes. "Other" = answer given but no theme keyword matched; "General / no specifics" = a bare yes/always/definitely; "No answer / unsure" = none, idk, not sure.',
 '• Q9 is coded as a status instead of themes: Yes / Some, but insufficient / Little or none / Unsure / Names an action or idea (the respondent suggested a solution rather than saying whether action is happening).',
 '• Answers that only name a topic number ("Topic 3", "iii") are tagged with that topic\'s theme (1 = Reproductive health & fertility, 2 = Access, cost & equity, 3 = Technology & AI).',
 '• Keyword tagging is approximate. Long answers can pick up extra themes from incidental words; spot-check the Tagged Responses sheet before quoting a specific number.']
for i, t_ in enumerate(lines):
    c = ns.cell(start + i, 1, t_)
    c.font = bold if i == 0 else body

wb.save(OUT)
print('saved', wb.sheetnames, 'summary rows', r)
