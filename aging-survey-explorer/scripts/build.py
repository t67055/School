import re, sys
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC, OUT = sys.argv[1], sys.argv[2]

# ---------------------------------------------------------------- helpers
def clean(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace('\xa0', ' ').strip()
    s = re.sub(r'[ \t]+', ' ', s)
    return s or None

def strip_letter(s):
    """Remove multiple-choice letter prefixes such as 'B. Female' or 'A. x, C. y'."""
    if s is None:
        return None
    return re.sub(r'(^|,\s*)[A-Z]\.\s+', r'\1', s).strip()

def load_respondents(ws, demo_cols, ncols, offset=0):
    """Read a sheet; rows with no demographic values are continuation lines of the
    respondent above (multi-line answers split across rows) and get folded in."""
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    out = []
    for i, r in enumerate(rows, start=2):
        r = [clean(c) for c in (list(r) + [None] * ncols)[:ncols]]
        if any(r[c] for c in demo_cols):
            out.append({'row': i, 'cells': r})
        elif out and any(r):
            prev = out[-1]['cells']
            for j, c in enumerate(r):
                if c:
                    prev[j] = f'{prev[j]}; {c}' if prev[j] else c
        # continuation rows before the first respondent are header overflow -> dropped
    return out

# ---------------------------------------------------------------- locations
US_STATES = {
 'AL':'Alabama','AK':'Alaska','AZ':'Arizona','AR':'Arkansas','CA':'California','CO':'Colorado',
 'CT':'Connecticut','DE':'Delaware','FL':'Florida','GA':'Georgia','HI':'Hawaii','ID':'Idaho',
 'IL':'Illinois','IN':'Indiana','IA':'Iowa','KS':'Kansas','KY':'Kentucky','LA':'Louisiana',
 'ME':'Maine','MD':'Maryland','MA':'Massachusetts','MI':'Michigan','MN':'Minnesota',
 'MS':'Mississippi','MO':'Missouri','MT':'Montana','NE':'Nebraska','NV':'Nevada',
 'NH':'New Hampshire','NJ':'New Jersey','NM':'New Mexico','NY':'New York','NC':'North Carolina',
 'ND':'North Dakota','OH':'Ohio','OK':'Oklahoma','OR':'Oregon','PA':'Pennsylvania',
 'RI':'Rhode Island','SC':'South Carolina','SD':'South Dakota','TN':'Tennessee','TX':'Texas',
 'UT':'Utah','VT':'Vermont','VA':'Virginia','WA':'Washington','WV':'West Virginia',
 'WI':'Wisconsin','WY':'Wyoming','DC':'District of Columbia'}
US = 'United States'
# name (lowercase) -> (state/province, country)
PLACES = {s.lower(): (s, US) for s in US_STATES.values()}
PLACES.update({
 'washington dc': ('District of Columbia', US), 'washington d.c.': ('District of Columbia', US),
 'new york city': ('New York', US), 'nyc': ('New York', US), 'baltimore': ('Maryland', US),
 'boston': ('Massachusetts', US), 'seattle': ('Washington', US), 'chicago': ('Illinois', US),
 'dallas': ('Texas', US), 'los angeles': ('California', US), 'westwood': ('California', US),
 'mayland': ('Maryland', US),
 # China
 'anhui': ('Anhui', 'China'), 'gansu': ('Gansu', 'China'), 'guangdong': ('Guangdong', 'China'),
 'jiangsu': ('Jiangsu', 'China'), 'jiangxi': ('Jiangxi', 'China'), 'sichuan': ('Sichuan', 'China'),
 'zhejiang': ('Zhejiang', 'China'), 'beijing': ('Beijing', 'China'), 'shanghai': ('Shanghai', 'China'),
 'chongqing': ('Chongqing', 'China'), 'suzhou': ('Jiangsu', 'China'), 'ningbo': ('Zhejiang', 'China'),
 'chengdu': ('Sichuan', 'China'), 'hong kong': ('Hong Kong', 'China'),
 # elsewhere
 'ontario': ('Ontario', 'Canada'), 'british columbia': ('British Columbia', 'Canada'),
 'andhra pradesh': ('Andhra Pradesh', 'India'), 'telangana': ('Telangana', 'India'),
 'melbourne': ('Victoria', 'Australia'), 'london': ('England', 'United Kingdom'),
 'cambridgeshire': ('England', 'United Kingdom'), 'england': ('England', 'United Kingdom'),
 'scotland': ('Scotland', 'United Kingdom'), 'athens': ('Attica', 'Greece'),
 'abidjan': ('Abidjan', "Côte d'Ivoire"), 'lusaka': ('Lusaka', 'Zambia'),
})
COUNTRIES = {
 'united states of america': US, 'united states': US, 'usa': US, 'us': US,
 'china': 'China', 'taiwan': 'Taiwan', 'chinese taipei': 'Taiwan', 'united kingdom': 'United Kingdom',
 'uk': 'United Kingdom', 'korea': 'South Korea', 'turkiye': 'Türkiye', 'gr': 'Greece',
 "côte d'ivoire": "Côte d'Ivoire",
}
for c in ['Australia','Bolivia','Brazil','Canada','El Salvador','Finland','France','Germany','Greece',
          'Hungary','India','Indonesia','Iran','Japan','Malaysia','Mexico','Myanmar','Nepal','Nigeria',
          'Peru','Philippines','Rwanda','Singapore','Switzerland','Ukraine','United Arab Emirates',
          'Venezuela','Zambia']:
    COUNTRIES[c.lower()] = c
ABBR = {k.lower(): ('state', v) for k, v in US_STATES.items()}
ABBR['bc'] = ('place', ('British Columbia', 'Canada'))
LOC_OVERRIDE = {
 'republic of china (taiwan)': ([], 'Taiwan'),
 'california, usa (or i guess baltimore?)': (['California'], US),
 'md/nj': (['Maryland', 'New Jersey'], US),
 'asia': ([], None),
}

def parse_location(raw):
    if not raw:
        return None, None, None
    s = raw.lower().strip()
    if s in LOC_OVERRIDE:
        st, co = LOC_OVERRIDE[s]
        note = 'location too vague to place' if s == 'asia' else None
        return ' / '.join(st) or None, co, note
    s = s.replace('p.r.china', 'china').replace('u.s.a.', 'usa').replace('u.s.', 'us')
    s = re.sub(r'\((jhu)\)', ' ', s)
    states, countries = [], []
    names = sorted(set(PLACES) | set(COUNTRIES), key=len, reverse=True)
    for n in names:
        pat = r'(?<![a-z])' + re.escape(n) + r'(?![a-z])'
        if re.search(pat, s):
            if n in PLACES:
                st, co = PLACES[n]
                if st not in states: states.append(st)
                countries.append(co)
            else:
                countries.append(COUNTRIES[n])
            s = re.sub(pat, ' ', s)
    for tok in re.findall(r'[a-z]+', s):
        if tok in ABBR:
            kind, val = ABBR[tok]
            if kind == 'state':
                if val not in states: states.append(val)
                countries.append(US)
            else:
                states.append(val[0]); countries.append(val[1])
        elif tok == 'ca':
            states.append('California'); countries.append(US)
        else:
            return None, None, f'unparsed token "{tok}"'
    countries = list(dict.fromkeys(countries))
    if len(countries) > 1:
        return ' / '.join(states) or None, ' / '.join(countries), 'multiple countries'
    return (' / '.join(states) or None), (countries[0] if countries else None), None

def combo(st, co):
    return ', '.join(p for p in (st, co) if p) or None

# ---------------------------------------------------------------- field normalizers
def norm_gender(g):
    if not g: return None
    g0 = strip_letter(g).lower()
    if g0 in ('male', 'm', 'man'): return 'Male'
    if g0 in ('female', 'f', 'woman', 'femail', 'feminine'): return 'Female'
    if 'non-binary' in g0 or 'genderqueer' in g0 or 'gender diverse' in g0: return 'Non-binary / gender diverse'
    if 'prefer not' in g0: return 'Prefer not to answer'
    if 'self-describe' in g0: return 'Self-described'
    return g

def norm_age(a):
    if not a: return None
    m = re.search(r'\d+(\.\d+)?', a)
    return int(float(m.group())) if m else None

BIN_LABELS = ['15% or less', '16–20%', '21–30%', 'Over 30%', 'Not sure']
def bin_pct(p):
    if p is None: return None
    if p <= 15: return BIN_LABELS[0]
    if p <= 20: return BIN_LABELS[1]
    if p <= 30: return BIN_LABELS[2]
    return BIN_LABELS[3]

def norm_pct_numeric(raw):
    """Returns (numeric percent or None, bin, note)."""
    if not raw: return None, None, None
    s = raw.replace('%', ' ').lower()
    m = re.fullmatch(r'\s*(\d+)\s*/\s*(\d+)\s*', s)
    if m:
        v = 100 * int(m.group(1)) / int(m.group(2))
        return round(v, 2), bin_pct(v), None
    nums = [float(n) for n in re.findall(r'\d+(?:\.\d+)?', s)]
    if not nums:
        if re.search(r"not sure|don.?t know|idk", s): return None, 'Not sure', None
        return None, None, f'Q6 answer not numeric: "{raw}"'
    if len(nums) == 2:
        v = sum(nums) / 2
        return v, bin_pct(v), f'Q6 range "{raw}" stored as midpoint'
    v = nums[0]
    if v < 1:  # answered as a fraction (0.25 = 25%)
        v = round(v * 100, 2)
    return v, bin_pct(v), None

def norm_pct_bin(raw):
    """Multiple-choice range answers (Groups 4, 6) -> common bins."""
    if not raw: return None
    s = strip_letter(raw).replace('–', '-').lower()
    if 'not sure' in s: return 'Not sure'
    if 'more than 30' in s or '>30' in s: return 'Over 30%'
    m = re.match(r'(\d+)-(\d+)', s)
    lo, hi = int(m.group(1)), int(m.group(2))
    return bin_pct(hi)

TOPIC = {1: 'Reproductive health & longevity', 2: 'Democratizing longevity (access)', 3: 'AI/ML for aging & health'}
def topic_of(s):
    if not s: return None
    t = s.lower()
    if 'equally' in t: return 'All three equally'
    if 'none of them' in t: return 'None of them'
    if 'reproductive' in t: return TOPIC[1]
    if 'democratiz' in t or 'accessible longevity' in t: return TOPIC[2]
    if 'machine learning' in t or re.search(r'\bai\b', t): return TOPIC[3]
    return s

def topic_num(s):
    t = topic_of(s)
    return {v: k for k, v in TOPIC.items()}.get(t)

YES = r'^(yes|yeah|yep|yea|definitely|absolutely|of course|for sure|always|certainly|sure|100%|1000%|i think so|it definitely)'
NO = r'^(no\b|nope|not really|nothing)'
UNSURE = r'^(maybe|perhaps|probably|possibly|somewhat|not sure|unsure|idk|i don.?t know|don.?t know|i.?m not sure|on some level|i suppose|kind of)'
def yes_no(raw):
    """Returns (Yes/No/Unsure or None, comment-or-None)."""
    if not raw: return None, None
    s = strip_letter(raw)
    l = s.lower().strip()
    if re.match(YES, l): v = 'Yes'
    elif re.match(NO, l): v = 'No'
    elif re.match(UNSURE, l): v = 'Unsure'
    else: v = None
    bare = re.sub(r'[^a-z ]', '', l).strip() in ('yes', 'no', 'not sure', 'maybe', 'yes for sure', 'definitely', 'absolutely', 'of course')
    return v, (None if bare else s)

def rating(raw):
    if not raw: return None
    try:
        v = float(raw)
        return int(v) if 1 <= v <= 5 else None
    except ValueError:
        return None

# ---- ranks
def rank_value(raw):
    """Per-topic rank cell -> 1/2/3."""
    if not raw: return None
    l = raw.lower()
    if 'most important' in l and not re.search(r'[23]', l): return 1
    if 'least important' in l and not re.search(r'[12]', l): return 3
    if l.strip() == 'medium': return 2
    m = re.search(r'[123]', l)
    return int(m.group()) if m else None

ROMAN = {'i': 1, 'ii': 2, 'iii': 3, 'io': 2}
ORDER_OVERRIDE = {
 'number 2 offers the most potential in terms of helping the population overall. research in number 1 might lead to some kind of breakthrough. i have doubts about number 3.': [2, 1, 3],
 '(iii) most important; (ii); (i) least important': [3, 2, 1],
 '#1 (ii) #2 (iii) #3 (i)': [2, 3, 1],
 '1. democratizing longevity science/solutions; 2. ai; 3. boosting reproduction.': [2, 3, 1],
 'sentence two, sentence, three, sentence one': [2, 3, 1],
}
SAME = ('same as prior rank', 'same', 'my hope is the exact order as above.')
def parse_order(raw):
    """Ordered list (most -> least important) -> ranks dict {topic: rank} or None."""
    if not raw: return None, None
    l = raw.lower().strip()
    l = re.sub(r'\s*\n\s*', '; ', l)
    if l in ORDER_OVERRIDE:
        order = ORDER_OVERRIDE[l]
    elif l in SAME:
        return 'SAME', None
    else:
        romans = re.findall(r'(?<![a-z])(iii|ii|io|i)(?![a-z])', l)
        words = []
        for kw, n in [('reproduct', 1), ('democra', 2), ('ai', 3), ('machine', 3)]:
            m = re.search(r'(?<![a-z])' + kw, l)
            if m: words.append((m.start(), n))
        digits = re.findall(r'[0-9]', l)
        if len(romans) == 3:
            order = [ROMAN[r] for r in romans]
        elif len(words) >= 3:
            order = [n for _, n in sorted(words)][:3]
        elif len(digits) == 3:
            order = [int(d) for d in digits]
        else:
            return None, f'rank answer not interpretable: "{raw}"'
    if sorted(order) != [1, 2, 3]:
        return None, f'rank answer not a valid 1-2-3 order: "{raw}"'
    return {t: pos + 1 for pos, t in enumerate(order)}, None

def parse_positional(cells):
    """Group 4: [most, next, least] cells each naming a topic."""
    ranks, note = {}, None
    picks = [[topic_num(p) for p in (c or '').split(', ') if p] for c in cells]
    for pos, p in enumerate(picks, start=1):
        if len(p) == 1 and p[0] and sum(p[0] in q for q in picks) == 1:
            ranks[p[0]] = pos
    if len(ranks) < 3:
        note = 'rank grid answered inconsistently; only unambiguous ranks kept'
    return ranks, note

def interviewer_map(wb):
    """First names / full names of group members -> full name, read from the Member_Groups sheet if present."""
    m = {}
    if 'Member_Groups' not in wb.sheetnames:
        return m
    for row in wb['Member_Groups'].iter_rows(min_row=2, values_only=True):
        for v in row[1:]:
            if isinstance(v, str) and v.strip():
                full = v.strip()
                m[full.lower()] = full
                m.setdefault(full.split()[0].lower(), full)
    return m

def join_topics(a, b, c):
    parts = [f'({n}) {v}' for n, v in zip(('i', 'ii', 'iii'), (a, b, c)) if v]
    return '; '.join(parts) or None

# ---------------------------------------------------------------- per-group column maps
# fields: age gender from res aging pct pctbin challenge short why action maintain cando
#         desc pr1-3 / porder / ppos ; sr1-3 / sorder / spos ; reason excite_topic excite concern_topic concern
#         ethics ethics_detail excited excited_r bio bio_r interviewer timestamp
MAPS = {
 'Group 1': dict(n=23, demo=[0,1,2,3], age=0, gender=1, frm=2, res=3, aging=4, pct=5, challenge=6, short=7,
                 action=8, maintain=9, cando=10, desc=11, pr=[12,13,14], sr=[15,16,17], reason=18,
                 excite=19, ethics=20, excited=21, bio=22),
 'Group 2': dict(n=21, demo=[0,1,2,3], age=0, gender=1, frm=2, res=3, aging=4, pct=5, challenge=6, short=7,
                 action=8, maintain=9, cando=10, desc=11, porder=12, sorder=13, reason=14,
                 excite_topic=15, concern_topic=16, ethics=17, ethics_detail=18, excited=19, bio=20),
 'Group 3': dict(n=24, demo=[0,1,2,3], age=0, gender=1, frm=2, res=3, aging=4, pct=5, challenge=6, short=7,
                 action=8, maintain=9, cando=10, desc=11, pr=[12,13,14], reason=15, sr=[16,17,18],
                 excite=19, concern=20, ethics=21, excited_r=22, bio_r=23),
 'Group 4': dict(n=29, demo=[1,2,3,4], timestamp=0, age=1, frm=2, res=3, gender=4, pctbin=5, aging=6,
                 challenge=7, short=8, why=9, action=10, maintain=11, cando=12, desc3=[13,14,15],
                 ppos=[16,17,18], reason_i=19, spos=[20,21,22], reason_s=23, excite_topic=24, concern=25,
                 ethics=26, excited_r=27, bio_r=28),
 'Group 5': dict(n=21, demo=[1,2,3,4], age=1, gender=2, frm=3, res=4, aging=5, pct=6, challenge=7, short=8,
                 action=9, maintain=10, cando=11, desc=12, porder=13, sorder=14, reason=15,
                 excite=16, concern=17, ethics=18, excited=19, bio=20),
 'Group 6': dict(n=30, demo=[1,2,3,4], age=1, gender=2, frm=3, res=4, aging=5, pctbin=6, challenge=7,
                 short=8, why=9, action=10, maintain=11, cando=12, cando_b=13, desc3=[14,15,16],
                 pr=[17,18,19], sr=[20,21,22], reason=23, excite_topic=24, concern=25, ethics=26,
                 ethics_detail=27, excited_r=28, bio_r=29),
 'Group 7': dict(n=25, demo=[0,1,2,3], age=0, gender=1, frm=2, res=3, aging=4, pct=5, challenge=6, short=7,
                 action=8, maintain=9, cando=10, desc=11, pr=[12,13,14], sr=[15,16,17], reason=18,
                 excite=19, concern=20, ethics=21, excited_r=22, bio=23, interviewer=24),
 'Group 8': dict(n=30, demo=[0,1,2,3,4,5], age=0, gender=1, frm_country=2, frm_state=3, res_country=4,
                 res_state=5, aging=6, pct=7, challenge=8, short=9, action=10, maintain=11, cando=12,
                 desc3=[13,14,15], pr=[16,17,18], sr=[19,20,21], reason=22, excite_topic=23, excite=24,
                 concern_topic=25, concern=26, ethics=27, excited=28, bio=29),
}

COLUMNS = [
 ('Respondent ID', 11), ('Group', 9), ('Source row', 8), ('Interviewer', 16), ('Response timestamp', 18),
 ('Age', 6), ('Age category', 14), ('Gender', 14),
 ('From – State/Province', 18), ('From – Country', 16), ('From (State, Country)', 26),
 ('Residence – State/Province', 18), ('Residence – Country', 16), ('Residence (State, Country)', 26),
 ('Q5 How would you define aging?', 40),
 ('Q6 Est. % of US population 65+', 12), ('Q6 Estimate (range)', 12),
 ('Q7 Greatest challenge facing the aging population', 40),
 ('Q8 Challenge to solve in the short term', 40), ('Q8 Why prioritize it', 30),
 ('Q9 Action being taken on this challenge?', 30),
 ('Q10 Three things to maintain through life', 34), ('Q11 What you can do to maintain them', 34),
 ('Q12 Description of the three topics', 40),
 ('Personal rank – Reproductive health & longevity', 13), ('Personal rank – Democratizing longevity', 13),
 ('Personal rank – AI/ML for aging & health', 13),
 ('Societal rank – Reproductive health & longevity', 13), ('Societal rank – Democratizing longevity', 13),
 ('Societal rank – AI/ML for aging & health', 13),
 ('Reason for ranking', 40),
 ('Most exciting topic', 22), ('What excites you', 36), ('Most concerning topic', 22), ('Greatest concern', 36),
 ('Ethical considerations needed?', 12), ('Ethical considerations (details)', 36),
 ('Excited about future of aging research?', 12), ('Excitement rating (1–5)', 10), ('Excitement comment', 30),
 ('Bioengineers can help?', 12), ('Bioengineer role rating (1–5)', 10), ('Bioengineers comment', 30),
 ('Data notes', 36),
]
COLNAMES = [c for c, _ in COLUMNS]

wb = openpyxl.load_workbook(SRC)
INTERVIEWERS = interviewer_map(wb)
records, unparsed = [], []
for g, m in MAPS.items():
    gnum = int(g.split()[1])
    resp = load_respondents(wb[g], m['demo'], m['n'])
    for k, rr in enumerate(resp, start=1):
        c, notes = rr['cells'], []
        f = lambda key: c[m[key]] if key in m else None
        rec = dict.fromkeys(COLNAMES)
        rec.update({'Respondent ID': f'G{gnum}-{k:03d}', 'Group': gnum, 'Source row': rr['row']})
        if 'interviewer' in m and f('interviewer'):
            raw_iv = re.sub(r'\s*\(.*\)$', '', f('interviewer')).strip()
            rec['Interviewer'] = INTERVIEWERS.get(raw_iv.lower(), raw_iv)
        if 'timestamp' in m and f('timestamp'):
            rec['Response timestamp'] = pd.to_datetime(f('timestamp')).to_pydatetime()
        rec['Age'] = norm_age(f('age'))
        if f('age') and rec['Age'] is not None and not re.fullmatch(r'\d+(\.0)?', f('age')):
            notes.append(f'age given as "{f("age")}"')
        rec['Gender'] = norm_gender(f('gender'))
        # locations
        for side, key in (('From', 'frm'), ('Residence', 'res')):
            if key in m:
                raw = f(key)
                st, co, note = parse_location(raw)
                if note:
                    notes.append(f'{side.lower()}: "{raw}" – {note}')
                    if 'unparsed' in note: unparsed.append(raw)
            else:
                co_raw, st_raw = f(key + '_country'), f(key + '_state')
                co = parse_location(co_raw)[1] if co_raw else None
                st = st_raw if (st_raw and (co in (None, US))) else None
                if st and not co: co = US
            rec[f'{side} – State/Province'] = st
            rec[f'{side} – Country'] = co
            rec[f'{side} (State, Country)'] = combo(st, co)
        rec['Q5 How would you define aging?'] = strip_letter(f('aging')) if g == 'Group 6' else f('aging')
        if 'pct' in m:
            v, b, note = norm_pct_numeric(f('pct'))
            rec['Q6 Est. % of US population 65+'] = v
            rec['Q6 Estimate (range)'] = b
            if note: notes.append(note)
        else:
            rec['Q6 Estimate (range)'] = norm_pct_bin(f('pctbin'))
        sl = strip_letter if g == 'Group 6' else (lambda s: s)
        rec['Q7 Greatest challenge facing the aging population'] = sl(f('challenge'))
        rec['Q8 Challenge to solve in the short term'] = sl(f('short'))
        rec['Q8 Why prioritize it'] = sl(f('why'))
        rec['Q9 Action being taken on this challenge?'] = sl(f('action'))
        rec['Q10 Three things to maintain through life'] = sl(f('maintain'))
        cando = sl(f('cando'))
        if 'cando_b' in m and f('cando_b'):
            cando = f'{cando}; most helpful actions: {strip_letter(f("cando_b"))}'
        rec['Q11 What you can do to maintain them'] = cando
        if 'desc3' in m:
            rec['Q12 Description of the three topics'] = join_topics(*[sl(c[i]) for i in m['desc3']])
        else:
            rec['Q12 Description of the three topics'] = f('desc')
        # ranks
        for kind, label in (('p', 'Personal'), ('s', 'Societal')):
            ranks = {}
            if kind + 'r' in m:
                ranks = {t: rank_value(c[i]) for t, i in zip((1, 2, 3), m[kind + 'r'])}
                vals = [v for v in ranks.values() if v]
                if len(vals) == 3 and len(set(vals)) < 3:
                    notes.append(f'{label.lower()} ranks repeat a number (as answered)')
            elif kind + 'order' in m:
                ranks, note = parse_order(f(kind + 'order'))
                if ranks == 'SAME':
                    ranks = {t: rec[f'Personal rank – {n}'] for t, n in zip((1, 2, 3), ('Reproductive health & longevity', 'Democratizing longevity', 'AI/ML for aging & health'))}
                    notes.append('societal rank answered "same as personal"')
                if note: notes.append(f'{label.lower()} {note}')
                ranks = ranks or {}
            elif kind + 'pos' in m:
                ranks, note = parse_positional([c[i] for i in m[kind + 'pos']])
                if note: notes.append(f'{label.lower()} {note}')
            for t, n in zip((1, 2, 3), ('Reproductive health & longevity', 'Democratizing longevity', 'AI/ML for aging & health')):
                rec[f'{label} rank – {n}'] = ranks.get(t)
        if 'reason_i' in m:
            parts = [f'Individual: {f("reason_i")}' if f('reason_i') else None,
                     f'Society: {f("reason_s")}' if f('reason_s') else None]
            rec['Reason for ranking'] = '; '.join(p for p in parts if p) or None
        else:
            rec['Reason for ranking'] = sl(f('reason'))
        rec['Most exciting topic'] = topic_of(f('excite_topic'))
        rec['What excites you'] = f('excite')
        rec['Most concerning topic'] = topic_of(f('concern_topic'))
        rec['Greatest concern'] = sl(f('concern'))
        # ethics
        yn, com = yes_no(f('ethics'))
        if g == 'Group 6' and f('ethics'):
            e = strip_letter(f('ethics'))
            yn = {'Definitely yes': 'Yes', 'Not sure': 'Unsure', 'No special safeguards are necessary': 'No'}.get(e, 'Yes' if e.startswith('Yes') else None)
            com = None if e == 'Definitely yes' else e
        rec['Ethical considerations needed?'] = yn
        det = [p for p in (com, strip_letter(f('ethics_detail')) if 'ethics_detail' in m else None) if p]
        rec['Ethical considerations (details)'] = '; '.join(det) or None
        # excitement / bioengineers
        for key, ynk, rk, ck in (('excited', 'Excited about future of aging research?', 'Excitement rating (1–5)', 'Excitement comment'),
                                 ('bio', 'Bioengineers can help?', 'Bioengineer role rating (1–5)', 'Bioengineers comment')):
            if key in m:
                yn, com = yes_no(f(key))
                rec[ynk], rec[ck] = yn, com
            if key + '_r' in m:
                rec[rk] = rating(f(key + '_r'))
        rec['Data notes'] = '; '.join(notes) or None
        records.append(rec)

df = pd.DataFrame(records, columns=COLNAMES)
df.to_pickle('combined.pkl')
print('respondents', len(df), df.groupby('Group').size().to_dict())
print('unparsed locations', sorted(set(unparsed)))

# ---------------------------------------------------------------- write workbook
FONT = 'Arial'
hdr_fill = PatternFill('solid', fgColor='1F3864')
hdr_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
body_font = Font(name=FONT, size=10)
thin = Side(style='thin', color='BFBFBF')

ws = wb.create_sheet('Combined', 0)
ws.append(COLNAMES)
for rec in records:
    ws.append([rec[c] for c in COLNAMES])
n = len(records) + 1
age_col = get_column_letter(COLNAMES.index('Age') + 1)
cat_col = COLNAMES.index('Age category') + 1
for r in range(2, n + 1):
    a = f'{age_col}{r}'
    ws.cell(r, cat_col).value = (f'=IF({a}="","",IF({a}>=Notes!$B$5,"Elderly (65+)",'
                                 f'IF({a}<=Notes!$B$4,"Student","Middle-aged")))')
for j, (name, w) in enumerate(COLUMNS, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
    h = ws.cell(1, j)
    h.font, h.fill = hdr_font, hdr_fill
    h.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
for row in ws.iter_rows(min_row=2, max_row=n):
    for cell in row:
        cell.font = body_font
        cell.alignment = Alignment(vertical='top', wrap_text=False)
ts_col = COLNAMES.index('Response timestamp') + 1
for r in range(2, n + 1):
    ws.cell(r, ts_col).number_format = 'yyyy-mm-dd hh:mm'
ws.row_dimensions[1].height = 54
ws.freeze_panes = 'B2'
tab = Table(displayName='CombinedData', ref=f'A1:{get_column_letter(len(COLNAMES))}{n}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
ws.add_table(tab)

# Notes sheet
ns = wb.create_sheet('Notes', 1)
B = Font(name=FONT, bold=True, size=10)
rows = [
 ('Combined survey dataset – notes', None),
 (None, None),
 ('Age category cutoffs (edit the blue cells; the Combined sheet updates)', None),
 ('Student: age at or below', 29),
 ('Elderly: age at or above', 65),
 ('Middle-aged', 'everyone between the two cutoffs'),
 (None, None),
 ('Respondents by group and age category', None),
 ('Group', 'Student'),
]
for r in rows:
    ns.append(list(r))
ns['C9'], ns['D9'], ns['E9'], ns['F9'] = 'Middle-aged', 'Elderly (65+)', 'Age missing', 'Total'
start = 10
for i, gnum in enumerate(range(1, 9)):
    r = start + i
    ns.cell(r, 1, gnum)
    for j, cat in enumerate(['Student', 'Middle-aged', 'Elderly (65+)'], start=2):
        ns.cell(r, j, f'=COUNTIFS(Combined!$B:$B,$A{r},Combined!$G:$G,"{cat}")')
    ns.cell(r, 5, f'=COUNTIFS(Combined!$B:$B,$A{r},Combined!$F:$F,"")')
    ns.cell(r, 6, f'=COUNTIF(Combined!$B:$B,$A{r})')
tr = start + 8
ns.cell(tr, 1, 'Total')
for j in range(2, 7):
    L = get_column_letter(j)
    ns.cell(tr, j, f'=SUM({L}{start}:{L}{tr-1})')
notes_text = [
 'How the data was combined',
 '• One row per respondent from all 8 group sheets (original sheets are kept unchanged after this tab). "Source row" is the row in the original group sheet.',
 '• Multi-line answers that had spilled onto extra rows in the originals (rows with no age/gender/location) were folded back into the respondent above, joined with "; ".',
 '• Header overflow rows at the top of Group 2 (topic list and example) were dropped.',
 '• Questions were aligned by meaning; where a group asked a question as two parts (e.g., Q8 challenge + why), the second part fills the matching column or is appended with "; ".',
 '• Group 6 multiple-choice letter prefixes (e.g., "B. ") were removed.',
 'Locations',
 '• Each location is split into State/Province and Country, with a combined "State, Country" column. Abbreviations (MD, Ca, NYC…) were expanded; cities were mapped to their state/province (Baltimore → Maryland, Suzhou → Jiangsu, London → England).',
 '• Answers giving only a country (e.g., "China", "USA") have a blank state. "Asia" could not be placed. Hong Kong is recorded as Hong Kong, China; Taiwan / Chinese Taipei as Taiwan.',
 'Standardized answers',
 '• Q6: answers given as fractions (0.25) were converted to percent (25). Q6 Estimate (range) puts every group on common bins so Groups 4 and 6 (range choices) are comparable with the numeric answers.',
 '• Ranks: 1 = most important, 3 = least important, for every group. Group 2 and Group 5 typed an ordered list (e.g., "ii, i, iii" or "2,1,3"); both were read as the topics listed from most to least important. Group 3 "Most/Medium/Least important" became 1/2/3; Group 4 rank-position grids were inverted to per-topic ranks. Answers that were not a valid order are blank and flagged in Data notes. Repeated ranks in grid questions (Groups 1, 7, 8) are kept as answered and flagged.',
 '• Yes/No columns are derived from the answer\'s opening word (Yes / No / Unsure); any answer longer than a bare yes/no is kept in the matching comment/details column. Groups 3, 4, 6, 7 used 1–5 rating scales, which go in the rating columns.',
 '• Group 1 Q16 asked for one excitement and one fear together; those answers are in "What excites you".',
 '• Gender: Man/Woman/M/F and spelling variants mapped to Male/Female; non-binary and genderqueer grouped as "Non-binary / gender diverse".',
 '• Group 7 interviewer first names were completed to full names from the Member_Groups sheet.',
]
r0 = tr + 2
for i, t in enumerate(notes_text):
    c = ns.cell(r0 + i, 1, t)
    c.font = B if not t.startswith('•') else body_font
for row in ns.iter_rows(min_row=1, max_row=tr):
    for cell in row:
        cell.font = body_font
for addr in ('A1',):
    ns[addr].font = Font(name=FONT, bold=True, size=13)
for addr in ('A3', 'A8'):
    ns[addr].font = B
for j in range(1, 7):
    ns.cell(9, j).font = B
    ns.cell(tr, j).font = B
ns['B4'].font = ns['B5'].font = Font(name=FONT, size=10, color='0000FF')
ns['B4'].fill = ns['B5'].fill = PatternFill('solid', fgColor='FFF2CC')
ns.column_dimensions['A'].width = 40
for L in 'BCDEF':
    ns.column_dimensions[L].width = 14

# order: Combined, Notes, Member_Groups, then original group sheets
if 'Member_Groups' in wb.sheetnames:
    wb.move_sheet('Member_Groups', offset=-(wb.sheetnames.index('Member_Groups') - 2))
wb.active = 0
wb.save(OUT)
print('saved', OUT, wb.sheetnames)
